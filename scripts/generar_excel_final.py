#!/usr/bin/env python3
"""
Arma el Excel final de Mundo Ferretero a partir de las dos bases JSON:

  data/ferreterias_ig.json  -> hoja "Ferreterias"  (a quien seguir y escribirle)
  data/proveedores_ig.json  -> hoja "Proveedores"  (a quien ofrecerle pauta)

Salida: ferreterias_instagram.xlsx
"""

import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

# El orden de arranque es el de Mateo: Ituzaingo primero y despues se abre.
ORDEN_ZONA = [
    "Ituzaingo", "Castelar", "Haedo", "Moron", "Hurlingham", "Merlo",
    "San Antonio de Padua", "Ramos Mejia", "Ciudadela", "Caseros",
]


def leer(nombre):
    with open(os.path.join(RAIZ, "data", nombre), encoding="utf-8") as f:
        return json.load(f)


def prioridad(entrada):
    """Ordena por cercania a Ituzaingo, despues por provincia y nombre."""
    donde = entrada.get("donde", "")
    for i, zona in enumerate(ORDEN_ZONA):
        if zona.lower() in donde.lower():
            return (0, i, entrada.get("nombre", ""))
    prov = entrada.get("prov", "zz")
    return (1, 0, f"{prov}|{entrada.get('nombre', '')}")


def escribir(ws, headers, anchos, filas):
    ws.append(headers)
    for celda in ws[1]:
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for fila in filas:
        ws.append(fila)

    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def link(ig):
    return f"https://www.instagram.com/{ig.lstrip('@')}/"


def main():
    wb = Workbook()

    ferreterias = sorted(leer("ferreterias_ig.json"), key=prioridad)
    ws = wb.active
    ws.title = "Ferreterias"
    escribir(
        ws,
        ["Usuario IG", "Nombre", "Donde queda", "Provincia", "Link"],
        [34, 46, 52, 22, 46],
        [
            [e["ig"], e["nombre"], e["donde"], e.get("prov", ""), link(e["ig"])]
            for e in ferreterias
        ],
    )

    proveedores = sorted(
        leer("proveedores_ig.json"), key=lambda e: (e.get("rubro", ""), e.get("nombre", ""))
    )
    ws2 = wb.create_sheet("Proveedores (pauta)")
    escribir(
        ws2,
        ["Usuario IG", "Nombre", "Rubro", "Donde queda", "Info", "Link"],
        [32, 42, 38, 42, 68, 46],
        [
            [e["ig"], e["nombre"], e.get("rubro", ""), e.get("donde", ""), e.get("info", ""), link(e["ig"])]
            for e in proveedores
        ],
    )

    salida = os.path.join(RAIZ, "ferreterias_instagram.xlsx")
    wb.save(salida)
    print(f"OK -> {salida}")
    print(f"  Ferreterias: {len(ferreterias)}")
    print(f"  Proveedores: {len(proveedores)}")


if __name__ == "__main__":
    main()
