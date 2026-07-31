#!/usr/bin/env python3
"""
Genera el Excel de ferreterias con Instagram para el outreach de Mundo Ferretero.

Las cuentas de la hoja "Ferreterias" salieron de busquedas web y cada una tiene
una URL de instagram real como fuente. Las de la hoja "A verificar" aparecieron
en las busquedas pero no se pudo confirmar que sean de Argentina / zona oeste.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# (usuario_ig, nombre, donde_queda)
FERRETERIAS = [
    # --- Ituzaingo (base) ---
    ("@barracaituzaingo", "Barraca Ituzaingo / Ferreteria Ituzaingo", "Ituzaingo - Av. Ituzaingo e/ Elias Regules y Plazoleta Martin Fierro"),
    ("@h30ferreteria.ituzaingo", "H30 Ferreteria Express", "Ituzaingo"),
    ("@nhmconstruccion", "NHM Construccion - Corralon y Ferreteria", "Ituzaingo"),
    ("@delferroviariomateriales", "Del Ferroviario Materiales", "Ituzaingo - Brandsen 1365"),

    # --- Castelar ---
    ("@ferreteriavictoria_", "Ferreteria Victoria", "Castelar"),
    ("@ferreteria_hidalgo.castelar", "Ferreteria Hidalgo", "Castelar"),

    # --- Haedo ---
    ("@ferreteria_haedo", "Ferreteria Haedo", "Haedo"),
    ("@ferreteriamele", "Ferreteria MELE", "Haedo"),
    ("@ferrecasamenendez", "Casa Menendez", "Haedo"),
    ("@ferreteriaradielec", "Radielec (Daniel Digangi)", "Haedo"),

    # --- Moron ---
    ("@ferreteriatotal", "Ferreteria Total Moron", "Moron centro"),
    ("@laferreteriaypintureria", "La Ferreteria y Pintureria", "Moron - Belgrano 497"),
    ("@ferreteriasolano", "Ferreteria Solano", "Moron"),
    ("@ferreteriadaleo", "Ferreteria Daleo", "Moron"),

    # --- Hurlingham ---
    ("@ferre.elcosito", "Ferreteria El Cosito", "Hurlingham - Av. Vergara 4124"),
    ("@ferreteria.roga", "Ferreteria ROGA", "Hurlingham"),
    ("@ferreterialaplaza", "Ferreteria La Plaza", "Hurlingham - Av. Roca 1205"),
    ("@corralonelportugues", "El Portugues - Ferreteria y +", "Hurlingham"),

    # --- Merlo ---
    ("@ferreteria.merlo", "Ferreteria Merlo", "Merlo"),
    ("@ferreteriamarmat", "Ferreteria Marmat (aberturas)", "Merlo"),
    ("@promeconsrl", "Promecon SRL - Ferreteria industrial", "Merlo - Constitucion 805 / Ruta 200"),
    ("@bulmaq_srl", "Bulmaq SRL - Ferreteria industrial", "Merlo"),

    # --- San Antonio de Padua ---
    ("@ferrepadua_", "Ferreteria Padua C.A.", "San Antonio de Padua"),
    ("@ferreteriaoscari", "Ferreteria Oscar I", "San Antonio de Padua - Noguera 540"),

    # --- Ramos Mejia ---
    ("@ferreteriaramosmejia", "Ferreteria Ramos Mejia", "Ramos Mejia"),
    ("@ferreteriaybuloneraramosmejia", "Ferreteria y Bulonera Ramos Mejia", "Ramos Mejia"),

    # --- Ciudadela / Caseros / Tres de Febrero ---
    ("@ferremoya", "Ferreteria Casa Moya", "Ciudadela"),
    ("@zanotti.ferreteria", "Zanotti Ferreteria", "Caseros"),
    ("@ferreteria.pablo", "Ferreteria y Cerrajeria Pablo", "Tres de Febrero - Tres de Febrero 3010"),

    # --- La Matanza ---
    ("@ferreteriasanjusto", "Ferreteria San Justo", "San Justo"),
    ("@ferreteriaprimerano", "Ferreteria Primerano", "San Justo"),
    ("@symeg_ferreteria_industrial", "Symeg - Ferreteria Industrial / Bulonera", "La Tablada - Av. Crovara 3516"),
    ("@ferreteria.el15", "Ferreteria El 15 - Sanitarios y Electricidad", "Villa Luzuriaga"),

    # --- San Miguel / Bella Vista / Jose C. Paz ---
    ("@ferreteriaboulevardsm", "Ferreteria Boulevard", "Bella Vista (San Miguel)"),
    ("@ferreteriaargentina16", "Ferreteria Argentina", "Jose C. Paz - Av. H. Yrigoyen 4236 y 2279"),

    # --- General Rodriguez / Lujan ---
    ("@flaslatas_gr", "Ferreteria Las Latas", "General Rodriguez"),
    ("@laferreterialujan", "La Ferreteria Lujan", "Lujan"),

    # --- CABA oeste ---
    ("@ferreteria_cerrajeriaescalada", "Ferreteria-Cerrajeria Escalada", "CABA - Villa Luro, Escalada 585"),
    ("@ferreteria_del_parque", "Ferreteria del Parque", "CABA - Villa del Parque, Nogoya 3402"),
    ("@ferreteriabuenosaires", "Ferreteria Buenos Aires", "CABA"),

    # --- Mayoristas / industriales que venden a ferreterias ---
    ("@zonamayorista.ar", "Zona Mayorista - Bulonera profesional", "Buenos Aires (mayorista)"),
    ("@ferrelaindustrial", "Ferreteria La Industrial", "Buenos Aires"),
    ("@centrodemateriales", "Centro de Materiales", "Buenos Aires"),
    ("@deltaferreteriaindustrial", "Delta Ferreteria Industrial", "Buenos Aires"),
    ("@laindustrial.sa", "La Industrial - Ferreteria y Buloneria", "Buenos Aires"),

    # --- Resto del pais (para cuando estires la zona) ---
    ("@ferreteriaituzaingo", "Ferreteria Ituzaingo", "Rosario, Santa Fe"),
    ("@casamartinferreteria", "Casa Martin - Ferreteria y Pintureria", "Recreo, Santa Fe"),
    ("@ferreteriaargentina", "Tu Ferreteria Argentina", "Villa Carlos Paz, Cordoba"),
    ("@ferremoreno149", "Ferreteria Moreno", "Parana, Entre Rios"),
    ("@liniers.ferreteria", "Liniers Ferreteria", "Mendoza"),
    ("@laferreteriadevilla", "La Ferreteria de Villa", "Cipolletti, Rio Negro"),
    ("@ferr.eteriasanantonio", "Ferreteria San Antonio", "Media Agua, San Juan"),
]

# Aparecieron en las busquedas pero no se pudo confirmar pais / localidad.
A_VERIFICAR = [
    ("@ferreteriafurlanetto", "Ferreteria Corralon Furlanetto", "salio buscando corralones Ituzaingo - confirmar localidad"),
    ("@ferreteriacorralonimi", "Ferreteria Corralon IMI", "confirmar localidad"),
    ("@corralonpucara1", "Pucara - Corralon y Ferreteria", "confirmar localidad"),
    ("@ferreteria_corralon.lar", "Ferreteria Corralon LAR (Luis Rodriguez)", "confirmar localidad"),
    ("@corralon.doncoco", "Don Coco - Corralon y Ferreteria", "confirmar localidad"),
    ("@ferreteria.corralon.sanantonio", "Corralon San Antonio", "confirmar localidad"),
    ("@ferreterianain", "Ferreteria Industrial NAIN", "confirmar localidad"),
    ("@shopferretero", "Shop Ferretero", "confirmar localidad"),
    ("@ferreterialargentina", "Ferreteria La Argentina S.A.", "confirmar localidad"),
    ("@morenomaterialessn", "Ferreteria Moreno / Moreno Materiales", "confirmar localidad"),
    ("@casamorenoferreterias", "Casa Moreno Ferreterias", "confirmar localidad"),
    ("@_ferreteriaramos_", "Ferreteria Ramos", "confirmar - hay varias homonimas"),
    ("@ferreteria.ramos10", "Ferreteria Ramos", "confirmar - hay varias homonimas"),
    ("@ferreteriasramos", "Ferreterias Ramos", "confirmar - hay varias homonimas"),
    ("@ferreteria_lapalmera", "Ferreteria La Palmera", "confirmar localidad"),
    ("@lapalmeraferreteria", "Ferreteria La Palmera", "confirmar localidad"),
    ("@ferreteriapalmar", "Ferreteria Palmar", "confirmar localidad"),
    ("@ferreteria.libertad", "Ferreteria Libertad", "confirmar - puede no ser Libertad (Merlo)"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)


def escribir_hoja(ws, filas, headers, ancho):
    ws.append(headers)
    for celda in ws[1]:
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for usuario, nombre, donde in filas:
        ws.append([usuario, nombre, donde, f"https://www.instagram.com/{usuario.lstrip('@')}/"])

    for i, w in enumerate(ancho, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:D{ws.max_row}"


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "Ferreterias"
    escribir_hoja(
        ws,
        FERRETERIAS,
        ["Usuario IG", "Nombre", "Donde queda", "Link"],
        [32, 42, 58, 46],
    )

    ws2 = wb.create_sheet("A verificar")
    escribir_hoja(
        ws2,
        A_VERIFICAR,
        ["Usuario IG", "Nombre", "Nota", "Link"],
        [32, 42, 58, 46],
    )

    wb.save("ferreterias_instagram.xlsx")
    print(f"OK - {len(FERRETERIAS)} ferreterias + {len(A_VERIFICAR)} a verificar")


if __name__ == "__main__":
    main()
